"""Builds the Matrix 1C workbook by filling the original template (template_1c.xlsx),
so the layout, fonts, borders and notes stay exactly like the official form."""
import io
import math
import os
import re
from copy import copy

import openpyxl
from openpyxl.styles import Alignment

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "template_1c.xlsx")
FIRST_DATA_ROW = 17   # row 16 holds the column numbers (-1 ... -6)
TEMPLATE_CAPACITY = 10  # rows 17..26 are blank in the template, NOTE starts at row 27
WIDTHS = {1: 26.7, 2: 26.1, 3: 26.1, 4: 26.1, 5: 26.1, 6: 27.9}
YEAR_CELLS = ("A5", "A6", "A7", "A14", "C15")


def _set_year(ws, year):
    # replace 2023 but protect ranges such as "MTDP 2023-2025"
    pat = re.compile(r"(?<![-\d])2023(?![-\d])")
    for ref in YEAR_CELLS:
        v = ws[ref].value
        if isinstance(v, str):
            ws[ref].value = pat.sub(str(year), v)


def _lines(text, width):
    if not text:
        return 1
    per_line = max(8, int(width * 0.95))
    return sum(max(1, math.ceil(len(p) / per_line)) for p in str(text).split("\n"))


def _safe_title(name, used):
    t = re.sub(r"[\\/*?:\[\]]", "-", str(name or "Matrix 1C")).strip()[:31] or "Matrix 1C"
    base, i = t, 2
    while t in used:
        suffix = f" ({i})"
        t = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(t)
    return t


def _fill(ws, meta, rows, year):
    _set_year(ws, year)
    ws["C9"].value = meta.get("sector") or None
    ws["C10"].value = meta.get("department") or None
    ws["C11"].value = meta.get("mfo") or None
    ws["C12"].value = meta.get("outcome") or None

    for r in (9, 10, 11, 12):
        ws.cell(r, 3)._style = copy(ws.cell(r, 1)._style)

    need = len(rows)
    if need > TEMPLATE_CAPACITY:
        ws.insert_rows(FIRST_DATA_ROW + TEMPLATE_CAPACITY, amount=need - TEMPLATE_CAPACITY)

    proto = {c: ws.cell(FIRST_DATA_ROW, c) for c in range(1, 7)}
    prev_result, group_start = object(), None
    for i, r in enumerate(rows):
        row = FIRST_DATA_ROW + i
        values = [r.get("result"), r.get("indicator"), r.get("target"), r.get("actual"),
                  r.get("issues"), r.get("reco")]
        same = (r.get("result") or "") == prev_result and r.get("result")
        if same:
            values[0] = None
        else:
            group_start = row
            prev_result = r.get("result") or ""
        for c in range(1, 7):
            cell = ws.cell(row, c)
            if i >= TEMPLATE_CAPACITY:   # inserted rows have no style: copy it from the first data row
                cell._style = copy(proto[c]._style)
            f = copy(cell.font)
            f.b = False
            cell.font = f
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.value = values[c - 1]
        longest = max(_lines(values[c - 1], WIDTHS[c]) for c in range(2, 7))
        if values[0]:
            longest = max(longest, _lines(values[0], WIDTHS[1]))
        ws.row_dimensions[row].height = max(15.6, longest * 13.2 + 3)

    # merge the "Results" cell down over rows that belong to the same result
    i = 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and (rows[j + 1].get("result") or "") == (rows[i].get("result") or "") and rows[i].get("result"):
            j += 1
        if j > i:
            ws.merge_cells(start_row=FIRST_DATA_ROW + i, start_column=1, end_row=FIRST_DATA_ROW + j, end_column=1)
        i = j + 1


def build_workbook(groups, year):
    """groups: list of {"meta": {...}, "rows": [{result, indicator, target, actual, issues, reco}]}.
    One sheet per department, all in the Matrix 1C layout. Returns xlsx bytes."""
    wb = openpyxl.load_workbook(TEMPLATE)
    base = wb["1C"]
    groups = groups or [{"meta": {}, "rows": []}]
    sheets = [base] + [wb.copy_worksheet(base) for _ in groups[1:]]
    used = set()
    for ws, g in zip(sheets, groups):
        ws.title = _safe_title(g["meta"].get("department"), used)
        _fill(ws, g["meta"], g["rows"], year)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
