"""Reading data that departments send in.

Supported:
  1. A filled-in Matrix 1C (.xlsx) - the official form itself.
  2. A simple input table (.xlsx / .csv) - see input_template_bytes().
"""
import io
import re

import openpyxl
import pandas as pd

LABELS = {"sector:": "sector", "department/office:": "department",
          "major final output:": "mfo", "organizational outcome:": "outcome"}

ALIASES = {
    "department": ["department", "office", "department/office", "dept"],
    "sector": ["sector"],
    "mfo": ["major final output", "mfo"],
    "outcome": ["organizational outcome", "outcome"],
    "result": ["result", "results", "development objective", "objective"],
    "indicator": ["performance indicator", "performance indicators", "indicator", "indicators"],
    "target": ["target", "target for indicators", "target for indicators / py"],
    "actual": ["actual", "actual accomplishment", "accomplishment"],
    "remarks": ["remarks", "remarks/challenges", "challenges", "issues", "issues/concerns/gaps",
                "issues/concerns", "concerns", "narrative"],
}

INPUT_COLUMNS = ["Department", "Sector", "Major Final Output", "Organizational Outcome", "Result",
                 "Performance Indicator", "Target", "Actual Accomplishment", "Remarks / Challenges"]


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_matrix_1c(file):
    """Read a filled Matrix 1C. Returns (meta, rows, year)."""
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb["1C"] if "1C" in wb.sheetnames else None
    if ws is None:
        for cand in wb.worksheets:
            if any(_s(c.value).lower() == "results" for row in cand.iter_rows(min_row=1, max_row=40, max_col=1) for c in row):
                ws = cand
                break
    if ws is None:
        raise ValueError("No Matrix 1C sheet found (expected a sheet named '1C').")

    meta, year, header = {}, None, None
    for r in range(1, 41):
        a = _s(ws.cell(r, 1).value)
        key = LABELS.get(a.lower())
        if key:
            for c in range(2, 7):
                v = _s(ws.cell(r, c).value)
                if v:
                    meta[key] = v
                    break
        m = re.search(r"PY\s*(20\d\d)", a)
        if m and not year:
            year = int(m.group(1))
        if a.lower() == "results" and header is None:
            header = r
    if header is None:
        raise ValueError("Could not find the 'Results' header row in the Matrix 1C sheet.")

    rows, last_result = [], ""
    for r in range(header + 1, ws.max_row + 1):
        a = _s(ws.cell(r, 1).value)
        if a.upper().startswith("NOTE"):
            break
        if re.fullmatch(r"-\d", a):          # the "(1) (2) ..." numbering row
            continue
        vals = [_s(ws.cell(r, c).value) for c in range(1, 7)]
        if a:
            last_result = a
        if not any(vals[1:5]):
            continue
        rows.append({"result": last_result, "indicator": vals[1], "target": vals[2], "actual": vals[3],
                     "remarks": vals[4], "issues_final": "", "reco_final": ""})
    return meta, rows, year


def read_table(file, filename):
    """Read the simple input table (xlsx or csv). Returns list of dict rows (department etc. per row)."""
    if filename.lower().endswith(".csv"):
        df = pd.read_csv(file, dtype=str).fillna("")
    else:
        df = pd.read_excel(file, dtype=str).fillna("")
    norm = {re.sub(r"\s*/\s*", "/", re.sub(r"\s+", " ", str(c).strip().lower())): c for c in df.columns}
    colmap = {}
    for key, names in ALIASES.items():
        for n in names:
            if n in norm:
                colmap[key] = norm[n]
                break
    if "indicator" not in colmap and "result" not in colmap:
        raise ValueError("Could not find the 'Result' or 'Performance Indicator' columns. "
                         "Please use the input template.")
    rows = []
    for _, r in df.iterrows():
        d = {k: _s(r[c]) for k, c in colmap.items()}
        if any(d.get(k) for k in ("result", "indicator", "target", "actual")):
            rows.append(d)
    return rows


def input_template_bytes():
    example = pd.DataFrame([{
        "Department": "MEO - Engineering", "Sector": "Infrastructure",
        "Major Final Output": "Roads and bridges constructed and maintained",
        "Organizational Outcome": "Improved road network and accessibility",
        "Result": "Improved farm-to-market road network",
        "Performance Indicator": "No. of kilometers of road concreted",
        "Target": "12", "Actual Accomplishment": "9",
        "Remarks / Challenges": "Failed bidding for two projects and delayed release of funds; backhoe under repair.",
    }], columns=INPUT_COLUMNS)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        example.to_excel(xw, sheet_name="Input", index=False)
        ws = xw.sheets["Input"]
        for i, w in enumerate([22, 20, 34, 34, 34, 34, 12, 20, 60], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        guide = pd.DataFrame({"How to fill in": [
            "One row per performance indicator. Delete the example row before sending.",
            "Target / Actual: use numbers (e.g. 12 and 9) or percentages (100% and 75%).",
            "Remarks / Challenges: write the reasons, problems or notes in your own words (English or Filipino). "
            "The system reads this text to produce the Issues/Concerns/Gaps and Recommendation columns.",
            "Department, Sector, MFO and Outcome may be repeated on every row.",
        ]})
        guide.to_excel(xw, sheet_name="Guide", index=False)
        xw.sheets["Guide"].column_dimensions["A"].width = 120
    return buf.getvalue()
